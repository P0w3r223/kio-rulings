"""Tryb pokazowy (ADR-0008 Z-1…Z-14) — korpus, atrapa, znaczniki i ścieżka końcowa.

Test końcowy jest bramką techniczną fazy 3 (ADR-0008 §10 pkt 1): `kio-tool demo` bez
`KIO_TOOL_CONTACT`, bez klucza, pod `--block-network` (reguła 20, zamek na gnieździe — pomiar 21)
przechodzi pobierz (tabela → zgoda) → przerwanie → wznów → eksport → szukaj.
"""

from __future__ import annotations

import json
import re
from functools import partial
from pathlib import Path

import httpx
import pytest
import yaml

from kio_tool import cli, config
from kio_tool.config import KATALOG_WYNIKOW, PLIK_BAZY
from kio_tool.demo import Pokaz, ZegarDemo, wczytaj_wzorce
from kio_tool.demo.atlas import AtlasPokazowy
from kio_tool.demo.korpus import (
    DOMENA,
    NUMER_DO,
    NUMER_OD,
    ODWOLUJACY,
    PROTOKOLANCI,
    PRZEWODNICZACY,
    ZAMAWIAJACY,
    ZNACZNIK_TRESCI,
    generuj,
)
from kio_tool.docid import SourceName, normalize_signature
from kio_tool.errors import StoreError
from kio_tool.exporter import ATRYBUCJA_POKAZU, PRZEDROSTEK_POKAZU
from kio_tool.httpclient import build_http_client
from kio_tool.source.contract import load_contract
from kio_tool.store import Store
from kio_tool.ui import texts
from kio_tool.ui.prompts import TAK, SkryptowyPrompter
from kio_tool.ui.texts import Pytanie
from tests.wsparcie_sondy import ZegarTestowy

KONTRAKT = load_contract(SourceName("atlas"))
KORPUS = generuj(wczytaj_wzorce())
ZLOTE = Path(__file__).resolve().parent / "examples" / "atlas"
WZORCE = Path(__file__).resolve().parent.parent / "kio_tool" / "demo" / "wzorce.yaml"


# --- korpus -------------------------------------------------------------------------------------


def test_korpus_jest_deterministyczny_i_ma_rozmiar_kwartalu() -> None:
    assert generuj(wczytaj_wzorce()) == KORPUS
    assert 300 <= len(KORPUS) <= 420


def test_sygnatury_pokazowe_leza_w_przedziale_e1_i_sa_poprawne() -> None:
    for d in KORPUS:
        for sygnatura in d.sygnatury:
            kanon = normalize_signature(sygnatura)
            assert kanon == sygnatura, "postać zmierzona — pole tożsamości zostaje wierne"
            assert NUMER_OD <= int(sygnatura.split()[1].split("/")[0]) <= NUMER_DO + 1


def test_adresy_tylko_w_domenie_zastrzezonej_i_znacznik_w_tresci() -> None:
    for d in KORPUS:
        assert d.url_zrodla.startswith(DOMENA + "/")
        assert d.tresc.splitlines()[0] == ZNACZNIK_TRESCI
        assert "uzp.gov.pl" not in d.tresc


def test_pula_osob_nie_wystepuje_w_prawdziwych_zlotych_plikach() -> None:
    """Fikcja z nazwami mówiącymi — i dowód, że żadna z nich nie jest nazwą z korpusu."""
    prawdziwe = " ".join(p.read_text(encoding="utf-8") for p in ZLOTE.glob("*.json"))
    for nazwa in (*PRZEWODNICZACY, *PROTOKOLANCI, *ODWOLUJACY, *ZAMAWIAJACY):
        assert nazwa not in prawdziwe


def test_wzorce_niosa_wylacznie_klucze_z_bialej_listy() -> None:
    """Z-11: `wzorce.yaml` bez pól osobowych — żadnego klucza spoza listy generatora."""
    dane = yaml.safe_load(WZORCE.read_text(encoding="utf-8"))
    assert set(dane) == {"zrodlo", "rozstrzygniecia", "rodzaje", "etykiety_przepisow", "ustawy_pzp"}
    assert re.fullmatch(r"[0-9a-f]{64}", dane["zrodlo"]["sha256_wejscia"])
    for wpis in dane["etykiety_przepisow"]:
        assert set(wpis) == {"etykieta", "liczba"} and wpis["etykieta"].startswith("art.")


# --- atrapa -------------------------------------------------------------------------------------


def _klient(atrapa: AtlasPokazowy) -> httpx.Client:
    return build_http_client(
        user_agent="test",
        allowed=frozenset({"atlasprzetargow.pl"}),
        transport=httpx.MockTransport(atrapa),
    )


def test_search_atrapy_dopasowuje_sygnature_a_nie_tresc() -> None:
    """Lustro „Pomiaru filtrów Atlasu" (2026-09-18): słowo z treści → 0, sygnatura → 1."""
    atrapa = AtlasPokazowy(KORPUS, KONTRAKT)
    p = KONTRAKT.parametry_listy
    adres = KONTRAKT.baza + KONTRAKT.punkty.lista
    with _klient(atrapa) as klient:
        po_tresci = klient.get(adres, params={p.filtry["fraza"]: "wadium"}).json()
        po_sygnaturze = klient.get(adres, params={p.filtry["fraza"]: KORPUS[0].sygnatury[0]}).json()
    licznik = KONTRAKT.ksztalt.lista.licznik
    assert po_tresci[licznik] == 0
    assert po_sygnaturze[licznik] >= 1


# --- znacznik bazy ------------------------------------------------------------------------------


def test_baza_pokazowa_odmawia_trybu_produkcyjnego_i_odwrotnie(tmp_path: Path) -> None:
    pokazowa = tmp_path / "pokaz.sqlite"
    with Store.open(pokazowa, clock=ZegarTestowy(), pokazowa=True) as s:
        assert s.pokazowa
    with pytest.raises(StoreError, match="trybu pokazowego"):
        Store.open(pokazowa, clock=ZegarTestowy())

    operatora = tmp_path / "korpus.sqlite"
    with Store.open(operatora, clock=ZegarTestowy()) as s:
        s.start_run(
            kanal="atlas",
            zakres="x",
            started_at="2024-01-01T00:00:00Z",
            kryteria="{}",
            fingerprint="x",
        )
    with pytest.raises(StoreError, match="nie jest bazą pokazową"):
        Store.open(operatora, clock=ZegarTestowy(), pokazowa=True)


# --- ścieżka końcowa ----------------------------------------------------------------------------


class _Operator(SkryptowyPrompter):
    """Skrypt odpowiedzi, który przy wyborze przebiegu do wznowienia bierze pierwszą pozycję."""

    def zapytaj(self, pytanie: Pytanie) -> str:
        if pytanie.tresc == texts.pytanie_wznowienia([]).tresc:
            self.zadane.append(pytanie)
            return pytanie.opcje[0].klucz
        return super().zapytaj(pytanie)


class _PrzerwijRaz:
    """Atrapa, w której operator wciska Ctrl+C przy trzydziestym dokumencie — raz."""

    def __init__(self, atrapa: AtlasPokazowy) -> None:
        self.atrapa = atrapa
        self.dokumentow = 0

    def __call__(self, zadanie: httpx.Request) -> httpx.Response:
        if zadanie.url.path != KONTRAKT.punkty.lista:
            self.dokumentow += 1
            if self.dokumentow == 30:
                raise KeyboardInterrupt
        return self.atrapa(zadanie)


def test_sciezka_pokazu_bez_kontaktu_i_bez_sieci(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("KIO_TOOL_CONTACT", raising=False)
    monkeypatch.setenv(KONTRAKT.tempo.klucz_api.zmienna, "krotki")  # Z-14: nie ma być czytany
    atrapa = AtlasPokazowy(KORPUS, KONTRAKT)
    pokaz = Pokaz(
        klient_factory=partial(
            build_http_client, transport=httpx.MockTransport(_PrzerwijRaz(atrapa))
        ),
        zegar=ZegarDemo(50),
        atrapa=atrapa,
    )
    operator = _Operator(
        [
            texts.MENU_POBIERZ,
            texts.CEL_DATY,
            "2024-01-01",
            "2024-01-31",
            TAK,
            texts.MENU_WZNOW,
            TAK,
            "jsonl",
            texts.MENU_SZUKAJ,
            "wadium",
            texts.WROC,
            texts.MENU_WYJDZ,
        ]
    )

    cli._uruchom_pokaz(od_nowa=True, prompter=operator, pokaz=pokaz)

    sciezka = config.katalog_pokazu() / PLIK_BAZY
    with Store.open(sciezka, clock=ZegarTestowy(), pokazowa=True) as store:
        styczen = sum(1 for d in KORPUS if d.data_wydania.month == 1)
        assert store.count("documents") == styczen
        assert store.list_runs(1)[0].status == "zakonczony"
    eksporty = list((config.katalog_pokazu() / KATALOG_WYNIKOW).glob("*.jsonl"))
    assert eksporty and all(p.name.startswith(PRZEDROSTEK_POKAZU) for p in eksporty)
    wiersz = json.loads(eksporty[0].read_text(encoding="utf-8").splitlines()[0])
    assert ATRYBUCJA_POKAZU in wiersz["cytowanie"]
    assert "Atlas Przetargów (https" not in wiersz["cytowanie"]


# --- znaczniki eksportu w każdym formacie (ADR-0008 §10 pkt 2) ----------------------------------


def _baza_pokazowa(tmp_path: Path, pokazowa: bool = True) -> Store:
    from kio_tool import pipeline
    from kio_tool.criteria import Criteria

    atrapa = AtlasPokazowy(KORPUS[:12], KONTRAKT)
    store = Store.open(tmp_path / "b.sqlite", clock=ZegarTestowy(), pokazowa=pokazowa)
    pipeline.pobierz(
        "atlas",
        Criteria.model_validate({"od": "2024-01-01", "do": "2024-01-31"}),
        store,
        zgoda=True,
        user_agent="test",
        klient_factory=partial(build_http_client, transport=httpx.MockTransport(atrapa)),
        klucz_z_srodowiska=False,
        zegar=ZegarTestowy(),
    )
    return store


@pytest.mark.parametrize("pokazowa", [True, False])
def test_znaczniki_pokazu_w_kazdym_formacie_i_w_zadnym_produkcyjnym(
    tmp_path: Path, pokazowa: bool
) -> None:
    from openpyxl import load_workbook

    from kio_tool import pipeline
    from kio_tool.exporter import FORMATY

    store = _baza_pokazowa(tmp_path, pokazowa)
    run_id = store.list_runs(1)[0].run_id
    wynik = pipeline.eksportuj(
        store, run_ids=(run_id,), formaty=FORMATY, out=tmp_path, zegar=ZegarTestowy()
    )
    for sciezka in wynik.sciezki:
        assert sciezka.name.startswith(PRZEDROSTEK_POKAZU) == pokazowa, sciezka.name
        pliki = sorted(sciezka.glob("*.md")) if sciezka.is_dir() else [sciezka]
        for plik in pliki:
            if plik.suffix == ".xlsx":
                wb = load_workbook(plik, read_only=True)
                meta = {r[0]: r[1] for r in wb["Metadane"].iter_rows(values_only=True)}
                assert meta["tryb"].startswith("POKAZOWY") == pokazowa
                continue
            tresc = plik.read_text(encoding="utf-8")
            if plik.name == "INDEX.md":
                # Spis katalogu nie jest rekordem — niesie znacznik jako wiersz `tryb`.
                assert ("tryb: POKAZOWY" in tresc) == pokazowa
                continue
            assert (ATRYBUCJA_POKAZU in tresc) == pokazowa, plik.name
