"""Granice trybu pokazowego (ADR-0008 Z-2, Z-3, Z-4, Z-13, Z-14) — czego `test_demo.py` nie mierzy.

`test_demo.py` niesie ścieżkę końcową i znaczniki w każdym formacie. Tu stoją krawędzie tych
samych zabezpieczeń: znacznik bazy w obie strony i przy każdej gałęzi odmowy, `--od-nowa`, które
kasuje wyłącznie bazę pokazową, pierwszy ekran i czas pokazu podpięte w `cli`, przedrostek `DEMO_`
pod `--out` z nazwą pliku, znacznik treści i pole `demo` w **każdym** rekordzie eksportu, zegar
pokazu i stronicowanie atrapy.
"""

from __future__ import annotations

import json
import shutil
import sqlite3
import time
from functools import partial
from pathlib import Path

import httpx
import pytest
import typer
from openpyxl import load_workbook
from typer.testing import CliRunner

from kio_tool import cli, config, pipeline
from kio_tool.config import PLIK_BAZY
from kio_tool.criteria import Criteria
from kio_tool.demo import (
    TEMPO_DOMYSLNE,
    TEMPO_ENV,
    TEMPO_MAKS,
    ZegarDemo,
    tempo_ze_srodowiska,
    wczytaj_wzorce,
    zbuduj_pokaz,
)
from kio_tool.demo.atlas import AtlasPokazowy
from kio_tool.demo.korpus import ZNACZNIK_TRESCI, DokumentPokazowy, generuj
from kio_tool.docid import SourceName
from kio_tool.errors import StoreError
from kio_tool.exporter import (
    ARKUSZ_METADANE,
    ARKUSZ_ORZECZENIA,
    ATRYBUCJA_POKAZU,
    PRZEDROSTEK_POKAZU,
)
from kio_tool.httpclient import build_http_client
from kio_tool.source.contract import load_contract
from kio_tool.store import ID_BAZY_POKAZOWEJ, Store
from kio_tool.ui import texts
from kio_tool.ui.prompts import NIE, SkryptowyPrompter
from kio_tool.ui.texts import Block
from tests.wsparcie_sondy import ZegarTestowy

KONTRAKT = load_contract(SourceName("atlas"))
KORPUS = generuj(wczytaj_wzorce())
MALY = KORPUS[:12]
"""Dwa dni robocze stycznia — dość na przebieg, eksport i odmowę, mało na czas suity."""
STYCZEN = Criteria.model_validate({"od": "2024-01-01", "do": "2024-01-31"})


class Ekran:
    """Widok `cli.view` podstawiony na czas testu — zbiera teksty zamiast rysować `rich`."""

    def __init__(self) -> None:
        self.teksty: list[str] = []
        self.bledy: list[str] = []

    def block(self, block: Block) -> None:
        self.teksty.append(block.as_text())

    def message(self, text: str) -> None:
        self.teksty.append(text)

    def warning(self, text: str) -> None:
        self.teksty.append(text)

    def error(self, text: str) -> None:
        self.bledy.append(text)


@pytest.fixture
def ekran(monkeypatch: pytest.MonkeyPatch) -> Ekran:
    ekran = Ekran()
    monkeypatch.setattr(cli, "view", ekran)
    return ekran


@pytest.fixture(autouse=True)
def _bez_poswiadczen(monkeypatch: pytest.MonkeyPatch) -> None:
    """Z-14: pokaz ma działać bez adresu kontaktowego; klucz za krótki dla `register_secret`
    wywróciłby przebieg, gdyby ktoś go przeczytał — to jest tu czujnik, nie dekoracja."""
    monkeypatch.delenv(config.CONTACT_ENV, raising=False)
    monkeypatch.setenv(KONTRAKT.tempo.klucz_api.zmienna, "krotki")
    monkeypatch.delenv(TEMPO_ENV, raising=False)


def _baza_pokazu() -> Path:
    return config.katalog_pokazu() / PLIK_BAZY


def _pobierz(store: Store, korpus: tuple[DokumentPokazowy, ...] = MALY) -> str:
    atrapa = AtlasPokazowy(korpus, KONTRAKT)
    wynik = pipeline.pobierz(
        "atlas",
        STYCZEN,
        store,
        zgoda=True,
        user_agent="test",
        klient_factory=partial(build_http_client, transport=httpx.MockTransport(atrapa)),
        klucz_z_srodowiska=False,
        zegar=ZegarTestowy(),
    )
    return wynik.run_id


def _baza_operatora(sciezka: Path) -> None:
    """Baza produkcyjna z jednym zakończonym przebiegiem i dokumentami."""
    with Store.open(sciezka, clock=ZegarTestowy()) as store:
        _pobierz(store)


# --- znacznik bazy: każda gałąź `Store._sprawdz_tryb` (Z-2) -------------------------------------


def test_znacznik_jedzie_z_plikiem_przy_kopiowaniu(tmp_path: Path) -> None:
    """Z-2: znacznik siedzi w pliku, nie w ścieżce — kopia pod neutralną nazwą dalej odmawia."""
    zrodlo = tmp_path / "pokaz.sqlite"
    with Store.open(zrodlo, clock=ZegarTestowy(), pokazowa=True) as store:
        _pobierz(store)
    kopia = tmp_path / "korpus.sqlite"
    shutil.copy(zrodlo, kopia)

    with pytest.raises(StoreError, match="trybu pokazowego"):
        Store.open(kopia, clock=ZegarTestowy())
    with sqlite3.connect(kopia) as polaczenie:
        assert polaczenie.execute("PRAGMA application_id").fetchone()[0] == ID_BAZY_POKAZOWEJ


def test_odmowa_otwarcia_zwalnia_plik_bazy(tmp_path: Path) -> None:
    """Odmowa w `__init__` zamyka połączenie — na Windowsie plik zajęty nie dałby się usunąć,
    a operator, który przeczytał „otwórz przez `kio-tool demo`", nie mógłby go nawet przenieść."""
    pokazowa = tmp_path / "pokaz.sqlite"
    with Store.open(pokazowa, clock=ZegarTestowy(), pokazowa=True):
        pass

    with pytest.raises(StoreError):
        Store.open(pokazowa, clock=ZegarTestowy())

    pokazowa.unlink()
    assert not pokazowa.exists()


def test_tryb_pokazowy_odmawia_bazy_z_dokumentami_bez_przebiegow(tmp_path: Path) -> None:
    """Druga połowa warunku „niepusta": dokumenty bez przebiegu (np. po ręcznym sprzątaniu
    `runs`) to nadal korpus operatora. Istniejący test mierzy wyłącznie przebieg bez dokumentów."""
    operatora = tmp_path / "korpus.sqlite"
    _baza_operatora(operatora)
    with sqlite3.connect(operatora) as polaczenie:
        polaczenie.execute("PRAGMA foreign_keys = OFF")
        for tabela in ("run_documents", "requests_log", "runs"):
            polaczenie.execute(f"DELETE FROM {tabela}")  # noqa: S608 — nazwy stałe z testu
    with Store.open(operatora, clock=ZegarTestowy()) as store:
        assert store.count("runs") == 0 and store.count("documents") > 0

    with pytest.raises(StoreError, match="nie jest bazą pokazową"):
        Store.open(operatora, clock=ZegarTestowy(), pokazowa=True)


def test_tryb_pokazowy_odmawia_pustej_bazy_z_cudzym_application_id(tmp_path: Path) -> None:
    """Obcy `application_id` to cudzy plik — pusty czy nie, pokaz go nie przejmuje."""
    obca = tmp_path / "obca.sqlite"
    with sqlite3.connect(obca) as polaczenie:
        polaczenie.execute("PRAGMA application_id = 1234")

    with pytest.raises(StoreError, match="nie jest bazą pokazową"):
        Store.open(obca, clock=ZegarTestowy(), pokazowa=True)
    with sqlite3.connect(obca) as polaczenie:
        assert polaczenie.execute("PRAGMA application_id").fetchone()[0] == 1234


def test_pusta_baza_dostaje_znacznik_i_od_tej_chwili_jest_pokazowa(tmp_path: Path) -> None:
    sciezka = tmp_path / "pusta.sqlite"
    with Store.open(sciezka, clock=ZegarTestowy()) as store:
        assert store.pokazowa is False

    with Store.open(sciezka, clock=ZegarTestowy(), pokazowa=True) as store:
        assert store.pokazowa is True
    with pytest.raises(StoreError, match="trybu pokazowego"):
        Store.open(sciezka, clock=ZegarTestowy())


# --- `kio-tool demo --od-nowa` (Z-2) ------------------------------------------------------------


def test_od_nowa_odmawia_skasowania_bazy_operatora_w_katalogu_pokazu(ekran: Ekran) -> None:
    """Baza operatora przeniesiona do katalogu pokazu: `--od-nowa` odmawia zdaniem i kodem, plik
    zostaje nietknięty, kreator się nie otwiera."""
    sciezka = _baza_pokazu()
    _baza_operatora(sciezka)
    with Store.open(sciezka, clock=ZegarTestowy()) as store:
        przed = (store.count("runs"), store.count("documents"))
    operator = SkryptowyPrompter([])

    with pytest.raises(typer.Exit) as wyjscie:
        cli._uruchom_pokaz(od_nowa=True, prompter=operator, pokaz=zbuduj_pokaz(korpus=MALY))

    assert wyjscie.value.exit_code != 0
    assert any("nie jest bazą pokazową" in b for b in ekran.bledy)
    assert operator.zadane == []
    with Store.open(sciezka, clock=ZegarTestowy()) as store:
        assert (store.count("runs"), store.count("documents")) == przed


@pytest.mark.parametrize(("od_nowa", "przebiegow"), [(True, 0), (False, 1)])
def test_od_nowa_kasuje_baze_pokazowa_a_bez_flagi_ja_zachowuje(
    ekran: Ekran, od_nowa: bool, przebiegow: int
) -> None:
    sciezka = _baza_pokazu()
    with Store.open(sciezka, clock=ZegarTestowy(), pokazowa=True) as store:
        _pobierz(store)

    cli._uruchom_pokaz(
        od_nowa=od_nowa,
        prompter=SkryptowyPrompter([texts.MENU_WYJDZ]),
        pokaz=zbuduj_pokaz(korpus=MALY),
    )

    with Store.open(sciezka, clock=ZegarTestowy(), pokazowa=True) as store:
        assert store.count("runs") == przebiegow


def test_od_nowa_zabiera_pliki_wal_i_shm_razem_z_baza() -> None:
    """Zostawiony `-wal` obcej sesji przykleiłby się do nowej bazy o tej samej nazwie."""
    sciezka = _baza_pokazu()
    with Store.open(sciezka, clock=ZegarTestowy(), pokazowa=True):
        pass
    wal = sciezka.with_name(sciezka.name + "-wal")
    shm = sciezka.with_name(sciezka.name + "-shm")
    wal.touch()
    shm.touch()

    cli._usun_baze_pokazowa(sciezka)

    assert not sciezka.exists() and not wal.exists() and not shm.exists()


def test_od_nowa_bez_bazy_nie_robi_nic() -> None:
    sciezka = _baza_pokazu()
    cli._usun_baze_pokazowa(sciezka)
    assert not sciezka.exists()


# --- `cli`: pierwszy ekran i czas pokazu (Z-3 znacznik 1, Z-13) ---------------------------------


def test_pokaz_otwiera_sie_ekranem_trybu_pokazowego(ekran: Ekran) -> None:
    cli._uruchom_pokaz(
        od_nowa=True,
        prompter=SkryptowyPrompter([texts.MENU_WYJDZ]),
        pokaz=zbuduj_pokaz(korpus=MALY),
    )

    assert ekran.teksty[0] == texts.pierwszy_ekran(pokaz=True).as_text()


def test_tabela_kosztow_pokazu_niesie_czas_produkcyjny_i_czas_pokazu(ekran: Ekran) -> None:
    """Z-13: tabela nie kłamie o produkcji — pokaz dopisuje swój czas obok. Odmowa po tabeli nie
    kosztuje żądania o dokument (zgoda z pytania, nie z flagi)."""
    pokaz = zbuduj_pokaz(korpus=MALY, tempo=10)
    operator = SkryptowyPrompter(
        [texts.MENU_POBIERZ, texts.CEL_DATY, "2024-01-01", "2024-01-31", NIE, texts.MENU_WYJDZ]
    )

    cli._uruchom_pokaz(od_nowa=True, prompter=operator, pokaz=pokaz)

    tabela = next(t for t in ekran.teksty if "czas w trybie pokazowym" in t)
    assert "Koszt przebiegu" in tabela
    with Store.open(_baza_pokazu(), clock=ZegarTestowy(), pokazowa=True) as store:
        assert store.count("documents") == 0


def test_polecenie_demo_jest_w_pomocy_z_flaga_od_nowa(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("TERM", "dumb")
    monkeypatch.setenv("COLUMNS", "200")
    wynik = CliRunner().invoke(cli.app, ["demo", "--help"])

    assert wynik.exit_code == 0, wynik.output
    assert "--od-nowa" in wynik.output


# --- eksport: znaczniki na każdym rekordzie i pod `--out` (Z-3, Z-4) ----------------------------


@pytest.fixture
def pokazowa(tmp_path: Path) -> Store:
    store = Store.open(tmp_path / "pokaz.sqlite", clock=ZegarTestowy(), pokazowa=True)
    _pobierz(store)
    return store


def _eksport(store: Store, out: Path, *formaty: str) -> tuple[Path, ...]:
    run_id = store.list_runs(1)[0].run_id
    wynik = pipeline.eksportuj(
        store, run_ids=(run_id,), formaty=formaty, out=out, zegar=ZegarTestowy()
    )
    return wynik.sciezki


def test_kazdy_wiersz_jsonl_pokazu_niesie_zdanie_o_fikcji_i_znacznik_rekordu(
    pokazowa: Store, tmp_path: Path
) -> None:
    """Z-3 (5) i (6) per rekord: wiersz JSONL wyjęty z pliku nie ma przedrostka nazwy, więc sam
    musi mówić, że jest fikcją — w cytowaniu, w atrybucji i w surowym rekordzie."""
    (sciezka,) = _eksport(pokazowa, tmp_path, "jsonl")
    wiersze = [json.loads(w) for w in sciezka.read_text(encoding="utf-8").splitlines()]
    pole_tresci = KONTRAKT.ksztalt.dokument.pole_tresci

    assert len(wiersze) == len(MALY)
    for wiersz in wiersze:
        assert wiersz["cytowanie"].startswith(ATRYBUCJA_POKAZU)
        assert wiersz["atrybucja"] == ATRYBUCJA_POKAZU
        assert wiersz["rekord"]["demo"] is True
        assert wiersz["rekord"][pole_tresci].splitlines()[0] == ZNACZNIK_TRESCI


def test_kazdy_wiersz_arkusza_pokazu_ma_atrybucje_pokazowa_a_tryb_jest_pierwszy(
    pokazowa: Store, tmp_path: Path
) -> None:
    (sciezka,) = _eksport(pokazowa, tmp_path, "xlsx")
    wb = load_workbook(sciezka, read_only=True)
    wiersze = list(wb[ARKUSZ_ORZECZENIA].iter_rows(values_only=True))
    kolumna = wiersze[0].index("atrybucja")
    meta = list(wb[ARKUSZ_METADANE].iter_rows(values_only=True))
    wb.close()

    assert [w[kolumna] for w in wiersze[1:]] == [ATRYBUCJA_POKAZU] * len(MALY)
    # Wiersz zerowy arkusza to nagłówek kolumn — `tryb` jest pierwszą **pozycją** metadanych.
    assert meta[1][0] == "tryb" and str(meta[1][1]).startswith("POKAZOWY")


@pytest.mark.parametrize(
    ("podana", "oczekiwana"),
    [
        # `--out` jest rdzeniem nazwy bez rozszerzenia (README, „Użycie") — stąd sam `styczen`.
        ("styczen", f"{PRZEDROSTEK_POKAZU}styczen.jsonl"),
        (f"{PRZEDROSTEK_POKAZU}styczen", f"{PRZEDROSTEK_POKAZU}styczen.jsonl"),
    ],
)
def test_out_z_nazwa_pliku_nie_zdejmuje_przedrostka_pokazu(
    pokazowa: Store, tmp_path: Path, podana: str, oczekiwana: str
) -> None:
    """Z-3 (3) „także pod `--out`" — i bez podwójnego przedrostka, gdy operator już go wpisał."""
    (sciezka,) = _eksport(pokazowa, tmp_path / "moje" / podana, "jsonl")

    assert sciezka.name == oczekiwana
    assert sciezka.parent == tmp_path / "moje"


def test_out_z_nazwa_pliku_w_bazie_produkcyjnej_zostaje_bez_przedrostka(tmp_path: Path) -> None:
    with Store.open(tmp_path / "korpus.sqlite", clock=ZegarTestowy()) as store:
        _pobierz(store)
        (sciezka,) = _eksport(store, tmp_path / "styczen", "jsonl")
        (xlsx,) = _eksport(store, tmp_path / "styczen", "xlsx")

    assert sciezka.name == "styczen.jsonl"
    wb = load_workbook(xlsx, read_only=True)
    meta = list(wb[ARKUSZ_METADANE].iter_rows(values_only=True))
    wb.close()
    assert meta[1] == ("tryb", "produkcyjny")


# --- zegar pokazu (Z-13) ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("podane", "tempo"), [(0.1, 1.0), (1.0, 1.0), (4.0, 4.0), (TEMPO_MAKS * 10, TEMPO_MAKS)]
)
def test_tempo_zegara_jest_przyciete_do_przedzialu(podane: float, tempo: float) -> None:
    """Tempo poniżej 1 spowalniałoby pokaz względem produkcji; ponad maksimum limiter tracił
    sens jako lekcja „tak wygląda odstęp"."""
    assert ZegarDemo(podane).tempo == tempo


def test_sen_pokazu_trwa_ulamek_a_zegar_monotoniczny_idzie_pelnym_krokiem(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`sleep` śpi `s / tempo`, ale `monotonic` przesuwa się o pełne `s` — limiter liczy odstępy
    produkcyjne. Ścienny zostaje prawdziwy: dziennik żądań ma daty, nie czas przyspieszony."""
    przespane: list[float] = []
    monkeypatch.setattr(time, "sleep", przespane.append)
    zegar = ZegarDemo(4.0)
    przed = zegar.monotonic()

    zegar.sleep(8.0)
    zegar.sleep(0)
    zegar.sleep(-1)

    assert przespane == [2.0]
    assert zegar.monotonic() - przed >= 8.0
    assert abs(zegar.wall() - time.time()) < 1.0


def test_czas_pokazu_to_czas_produkcyjny_podzielony_przez_tempo() -> None:
    assert zbuduj_pokaz(korpus=MALY, tempo=5).czas_pokazu(100.0) == pytest.approx(20.0)


@pytest.mark.parametrize(("wartosc", "tempo"), [("12.5", 12.5), ("szybko", TEMPO_DOMYSLNE)])
def test_tempo_ze_srodowiska_przyjmuje_liczbe_a_smiec_zastepuje_domyslnym(
    monkeypatch: pytest.MonkeyPatch, wartosc: str, tempo: float
) -> None:
    monkeypatch.setenv(TEMPO_ENV, wartosc)
    assert tempo_ze_srodowiska() == tempo


def test_tempo_bez_zmiennej_jest_domyslne() -> None:
    assert tempo_ze_srodowiska() == TEMPO_DOMYSLNE


# --- atrapa: stronicowanie, filtry, 404, znacznik rekordu ---------------------------------------


def _lista(atrapa: AtlasPokazowy, **params: str) -> dict[str, object]:
    with build_http_client(
        user_agent="test",
        allowed=frozenset({"atlasprzetargow.pl"}),
        transport=httpx.MockTransport(atrapa),
    ) as klient:
        wynik: dict[str, object] = klient.get(
            KONTRAKT.baza + KONTRAKT.punkty.lista, params=params
        ).json()
        return wynik


def test_strony_atrapy_skladaja_sie_w_komplet_bez_powtorzen() -> None:
    """Lista przechodzi przez adapter stronami; zgubiona albo zdublowana strona w atrapie uczyłaby
    wznawiania na liczbach, których Atlas nie zwraca."""
    atrapa = AtlasPokazowy(KORPUS, KONTRAKT)
    p, lista = KONTRAKT.parametry_listy, KONTRAKT.ksztalt.lista
    oczekiwane = {d.slug for d in KORPUS if d.data_wydania.month == 1}
    zebrane: list[str] = []
    strona = 1
    while True:
        odp = _lista(
            atrapa,
            **{p.od: "2024-01-01", p.do: "2024-01-31", p.na_strone: "7", p.strona: str(strona)},
        )
        assert odp[lista.licznik] == len(oczekiwane)
        rekordy = odp[lista.klucz]
        assert isinstance(rekordy, list)
        zebrane.extend(r[lista.rekord.referencja] for r in rekordy)
        if not odp[lista.ma_wiecej]:
            break
        strona += 1

    assert len(zebrane) == len(set(zebrane))
    assert set(zebrane) == oczekiwane
    assert strona == -(-len(oczekiwane) // 7)


def test_filtr_dat_atrapy_obejmuje_oba_brzegi() -> None:
    dzien = KORPUS[0].data_wydania.isoformat()
    p = KONTRAKT.parametry_listy
    odp = _lista(AtlasPokazowy(KORPUS, KONTRAKT), **{p.od: dzien, p.do: dzien})

    assert odp[KONTRAKT.ksztalt.lista.licznik] == sum(
        1 for d in KORPUS if d.data_wydania.isoformat() == dzien
    )


def test_filtr_rozstrzygniecia_atrapy_to_rownosc() -> None:
    p = KONTRAKT.parametry_listy
    odp = _lista(AtlasPokazowy(KORPUS, KONTRAKT), **{p.filtry["rozstrzygniecie"]: "umorzono"})

    oczekiwane = sum(1 for d in KORPUS if d.rozstrzygniecie == "umorzono")
    assert 0 < odp[KONTRAKT.ksztalt.lista.licznik] == oczekiwane  # type: ignore[operator]


def test_atrapa_odpowiada_404_na_nieznany_slug_i_nieznana_sciezke() -> None:
    atrapa = AtlasPokazowy(MALY, KONTRAKT)
    with build_http_client(
        user_agent="test",
        allowed=frozenset({"atlasprzetargow.pl"}),
        transport=httpx.MockTransport(atrapa),
    ) as klient:
        brak = klient.get(KONTRAKT.baza + KONTRAKT.punkty.dokument + "/kio-1-24")
        obca = klient.get(KONTRAKT.baza + "/api/cokolwiek")
        jest = klient.get(KONTRAKT.baza + KONTRAKT.punkty.dokument + "/" + MALY[0].slug)

    assert brak.status_code == 404 and obca.status_code == 404
    assert jest.status_code == 200
    assert atrapa.zadan == 3


def test_rekord_atrapy_niesie_pole_demo_i_znacznik_w_pierwszym_wierszu_tresci() -> None:
    """Z-3 (6) u źródła: znacznik jest w surowcu, więc przeżyje każdą drogę przez potok."""
    atrapa = AtlasPokazowy(MALY, KONTRAKT)
    for d in MALY:
        rekord = atrapa.rekord(d)
        tresc = rekord[KONTRAKT.ksztalt.dokument.pole_tresci]
        assert rekord["demo"] is True
        assert isinstance(tresc, str) and tresc.splitlines()[0] == ZNACZNIK_TRESCI
