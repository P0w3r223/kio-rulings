"""`exporter.py`: układ formatów, sanityzacja, `full_text` poza arkuszem, `thesis*` poza korpusem.

Materiałem jest złoty dokument pomiaru 3a z **dopisanymi** polami odrzuconymi o wymyślonej
treści — w prawdziwym dokumencie `thesis` jest `null`, a test reguły 19 potrzebuje pola, które
ma co przepuścić. Lista pól odrzuconych pochodzi z `contract.yaml` (`pola_odrzucone`), nie
z pamięci autora: pole dopisane tam przez pomiar bez wiersza w tym teście byłoby polem bez
strażnika, a `thesis_snippet` jest dokładnie takim polem — dopisanym po `thesis`.

Blok atrybucji (reguła 15) ma osobny plik `test_attribution.py`.
"""

from __future__ import annotations

import csv
import json
from collections.abc import Iterator
from pathlib import Path

import openpyxl
import pytest

from kio_tool import exporter
from kio_tool.docid import SourceName
from kio_tool.errors import ExportError
from kio_tool.exporter import FORMATY, NAZWY_KOLUMN, Wpis, eksportuj, wiersz
from kio_tool.parser.details import rekord_z_bajtow, wyczytaj
from kio_tool.pipeline import mapa_pol
from kio_tool.safetext import FORMULA_PREFIXES, sanitize_text
from kio_tool.source.contract import load_contract

ZLOTE = Path(__file__).resolve().parent / "examples" / "atlas"
DOKUMENT_BAJTY = (ZLOTE / "dokument_20260918T103526Z.json").read_bytes()
KONTRAKT = load_contract(SourceName("atlas"))
MAPA = mapa_pol(KONTRAKT)
TEZA = "TEZA WYMYSLONA NA POTRZEBY TESTU"
POLA_ODRZUCONE: tuple[str, ...] = tuple(p.pole for p in KONTRAKT.pola_odrzucone)
"""Pola, których kontrakt Atlasu nie wpuszcza do korpusu pochodnego (dziś `thesis`,
`thesis_snippet`) — czytane z pliku danych, żeby nie być drugą listą obok niego."""
ZNACZNIKI: dict[str, object] = {
    pole: f"WYMYSLONA TRESC POLA {pole.upper()}" for pole in POLA_ODRZUCONE
}
ATRYBUCJA = KONTRAKT.licencja.atrybucja
METADANE: tuple[tuple[str, object], ...] = (("kryteria", "test"), ("dokumentow_w_eksporcie", 1))


def wpis(**zmiany: object) -> Wpis:
    rekord = rekord_z_bajtow(DOKUMENT_BAJTY)
    rekord["thesis"] = TEZA
    rekord.update(zmiany)
    return Wpis(
        doc_id="atlas:kio-1205-20",
        source="atlas",
        source_ref="kio-1205-20",
        sha256="d" * 64,
        fetched_at="2026-09-18T12:00:00Z",
        szczegoly=wyczytaj(rekord, MAPA),
        rekord=rekord,
        atrybucja=ATRYBUCJA,
    )


def zrodlo(*wpisy: Wpis) -> exporter.Zrodlo:
    def _iter() -> Iterator[Wpis]:
        yield from wpisy

    return _iter


def arkusz(sciezka: Path, nazwa: str) -> list[tuple[object, ...]]:
    wb = openpyxl.load_workbook(sciezka, read_only=True)
    return [tuple(w) for w in wb[nazwa].iter_rows(values_only=True)]


@pytest.fixture
def eksport(tmp_path: Path) -> dict[str, Path]:
    sciezki = eksportuj(tmp_path / "proba", zrodlo(wpis()), formaty=FORMATY, metadane=METADANE)
    return dict(zip(FORMATY, sciezki, strict=True))


# --- układ --------------------------------------------------------------------------------------


def test_xlsx_ma_trzy_arkusze_a_orzeczenia_naglowek_ze_slownika(eksport: dict[str, Path]) -> None:
    wb = openpyxl.load_workbook(eksport["xlsx"], read_only=True)

    assert wb.sheetnames == ["Orzeczenia", "Slownik", "Metadane"]
    wiersze = arkusz(eksport["xlsx"], "Orzeczenia")
    assert wiersze[0] == NAZWY_KOLUMN
    assert len(wiersze) == 2
    slownik = arkusz(eksport["xlsx"], "Slownik")
    assert [w[0] for w in slownik[1:]] == list(NAZWY_KOLUMN)
    metadane = arkusz(eksport["xlsx"], "Metadane")
    assert ("kryteria", "test") in metadane and ("dokumentow_w_eksporcie", 1) in metadane


def test_wiersz_orzeczenia_niesie_sygnature_date_i_tozsamosc(eksport: dict[str, Path]) -> None:
    naglowek, dane = arkusz(eksport["xlsx"], "Orzeczenia")
    komorki = dict(zip(naglowek, dane, strict=True))

    assert komorki["sygnatura"] == "KIO 1205/20"
    assert komorki["data_wydania"] == "2004-01-29"
    assert komorki["doc_id"] == "atlas:kio-1205-20" and komorki["sha256"] == "d" * 64
    assert komorki["dlugosc_tresci"] == 3021
    assert komorki["koszty"] is None


def test_csv_ma_te_same_kolumny_co_arkusz_z_bom_i_srednikiem(eksport: dict[str, Path]) -> None:
    surowe = eksport["csv"].read_bytes()
    assert surowe.startswith(b"\xef\xbb\xbf")
    with eksport["csv"].open(encoding="utf-8-sig", newline="") as plik:
        wiersze = list(csv.reader(plik, delimiter=";"))

    assert wiersze[0] == list(NAZWY_KOLUMN)
    assert len(wiersze) == 2 and wiersze[1][0] == "KIO 1205/20"


def test_md_to_katalog_z_plikiem_na_dokument_i_indeksem(eksport: dict[str, Path]) -> None:
    katalog = eksport["md"]
    plik = katalog / "atlas_kio-1205-20.md"

    assert plik.is_file() and (katalog / "INDEX.md").is_file()
    tresc = plik.read_text(encoding="utf-8")
    assert tresc.startswith("---\n") and 'sygnatura: "KIO 1205/20"' in tresc
    assert "Sygn. akt: KIO 1205/20" in tresc, "pełny tekst jest w Markdown"
    indeks = (katalog / "INDEX.md").read_text(encoding="utf-8")
    assert "[atlas_kio-1205-20.md](atlas_kio-1205-20.md)" in indeks and "kryteria: test" in indeks


# --- reguła 19 i limit komórki ------------------------------------------------------------------


def test_full_text_nie_wchodzi_do_arkusza_ani_csv_a_jest_w_jsonl_i_md(
    eksport: dict[str, Path],
) -> None:
    """Komórka Excela mieści 32 767 znaków; orzeczenie bywa dłuższe — treść idzie gdzie indziej."""
    naglowek, dane = arkusz(eksport["xlsx"], "Orzeczenia")
    assert "full_text" not in naglowek and "tresc" not in naglowek
    assert not any(isinstance(k, str) and "Sygn. akt: KIO 1205/20" in k for k in dane)
    assert "Sygn. akt: KIO 1205/20" not in eksport["csv"].read_text(encoding="utf-8-sig")

    obiekt = json.loads(eksport["jsonl"].read_text(encoding="utf-8").splitlines()[0])
    assert obiekt["rekord"]["full_text"].startswith("Sygn. akt: KIO 1205/20")


def test_pola_odrzucone_z_kontraktu_nie_wchodza_do_orzeczen_csv_ani_md_a_zostaja_w_jsonl(
    tmp_path: Path,
) -> None:
    """Reguła 19: JSONL jest zrzutem surowca (rekord w całości), reszta korpusem pochodnym.

    Pola brane z `contract.yaml`, nie wpisane tutaj: `thesis_snippet` dopisano do kontraktu po
    `thesis` i przez chwilę miało strażnika wyłącznie strukturalnego (`Szczegoly` nie ma na nie
    pola). Strukturalny wystarcza, dopóki nikt nie doda pola — a dodanie pola do `Szczegoly` jest
    zmianą jednej linii, więc zachowanie eksportu ma własny obserwator.
    """
    assert len(POLA_ODRZUCONE) >= 2 and "thesis" in POLA_ODRZUCONE, (
        "kontrakt Atlasu deklaruje pola odrzucone — bez nich ten test nie mierzy niczego"
    )
    sciezki = dict(
        zip(
            FORMATY,
            eksportuj(tmp_path / "o", zrodlo(wpis(**ZNACZNIKI)), formaty=FORMATY, metadane=()),
            strict=True,
        )
    )
    naglowek, dane = arkusz(sciezki["xlsx"], "Orzeczenia")
    arkusz_tekst = " ".join(str(k) for k in (*naglowek, *dane))
    csv_tekst = sciezki["csv"].read_text(encoding="utf-8-sig")
    md_tekst = "".join(p.read_text(encoding="utf-8") for p in sciezki["md"].glob("*.md"))
    obiekt = json.loads(sciezki["jsonl"].read_text(encoding="utf-8").splitlines()[0])

    for pole, znacznik in ZNACZNIKI.items():
        assert pole not in " ".join(str(n) for n in naglowek), f"kolumna `{pole}` w arkuszu"
        for fmt, tekst in (("xlsx", arkusz_tekst), ("csv", csv_tekst), ("md", md_tekst)):
            assert str(znacznik) not in tekst, f"{fmt}: pole `{pole}` przeszło do eksportu"
        assert obiekt["rekord"][pole] == znacznik, (
            f"JSONL jest zrzutem surowca — pole `{pole}` ma w nim zostać w całości"
        )
    assert obiekt["doc_id"] == "atlas:kio-1205-20" and obiekt["sha256"] == "d" * 64


def test_wiersz_bierze_wartosci_wylacznie_ze_szczegolow() -> None:
    dane = wiersz(wpis())

    assert set(dane) == set(NAZWY_KOLUMN)
    assert dane["organ"] == exporter.ORGAN and dane["atrybucja"] == ATRYBUCJA


# --- dane wrogie ------------------------------------------------------------------------------


def test_komorka_zaczynajaca_sie_od_formuly_dostaje_apostrof_w_xlsx_i_csv(tmp_path: Path) -> None:
    wrogi = wpis(appellant_raw='=HYPERLINK("https://example.org")\x07', chairperson="‮Nowak")
    sciezki = eksportuj(tmp_path / "w", zrodlo(wrogi), formaty=("xlsx", "csv"), metadane=())

    naglowek, dane = arkusz(sciezki[0], "Orzeczenia")
    komorki = dict(zip(naglowek, dane, strict=True))
    assert komorki["odwolujacy"] == '\'=HYPERLINK("https://example.org")'
    assert komorki["przewodniczacy"] == "Nowak"
    with sciezki[1].open(encoding="utf-8-sig", newline="") as plik:
        wiersze = list(csv.reader(plik, delimiter=";"))
    assert wiersze[1][naglowek.index("odwolujacy")].startswith("'=")


PREFIKSY_WIDOCZNE = tuple(p for p in FORMULA_PREFIXES if p == p.strip())
"""Prefiksy formuły, które mogą stanąć na początku **pola metadanych** rekordu.

`parser.details._napis` obcina białe znaki każdego pola metadanych, więc tabulator nigdy nie
dojdzie tamtą drogą do komórki — to jest podział ścieżek, nie wybór wygodnych przypadków, i ma
niżej własny obserwator. Tabulator dochodzi drogą, która białych znaków nie obcina, i tam jest
sprawdzany.
"""


@pytest.mark.parametrize("prefiks", PREFIKSY_WIDOCZNE, ids=repr)
def test_prefiks_formuly_w_polu_kanalu_jest_neutralizowany_w_arkuszu_i_csv(
    tmp_path: Path, prefiks: str
) -> None:
    """Prefiksy brane z `safetext.FORMULA_PREFIXES`, a nie wypisane tutaj.

    Wersja z CEIDG miała na tej liście `\\r`, którego `strip_control` zdejmował wcześniej — czyli
    stała deklarowała pokrycie, którego nie miała, i nikt tego nie mierzył, bo test znał tylko
    `=`. Parametryzacja z produkcji sprawia, że prefiks dopisany do listy bez pokrycia zapala się
    sam, a usunięty przestaje być sprawdzany razem z nim.

    Minus jest tu ważniejszy od `=`: w polskim tekście prawniczym myślnik na początku wartości
    jest częsty (wyliczenia, zakresy), więc to on mówi, czy neutralizacja działa na treści
    z korpusu, a nie tylko na wejściu wymyślonym pod test.
    """
    wrogi = wpis(appellant_raw=f"{prefiks}1+1")
    sciezki = eksportuj(tmp_path / "f", zrodlo(wrogi), formaty=("xlsx", "csv"), metadane=())

    naglowek, dane = arkusz(sciezki[0], "Orzeczenia")
    komorki = dict(zip(naglowek, dane, strict=True))
    with sciezki[1].open(encoding="utf-8-sig", newline="") as plik:
        wiersze = list(csv.reader(plik, delimiter=";"))

    assert komorki["odwolujacy"] == f"'{prefiks}1+1", "arkusz przepuścił prefiks formuły"
    assert wiersze[1][naglowek.index("odwolujacy")] == f"'{prefiks}1+1", "CSV przepuścił prefiks"


@pytest.mark.parametrize("prefiks", FORMULA_PREFIXES, ids=repr)
def test_prefiks_formuly_w_arkuszu_metadanych_jest_neutralizowany(
    tmp_path: Path, prefiks: str
) -> None:
    """Arkusz `Metadane` niesie `cel_pobrania`, czyli zdanie wpisane przez operatora flagą
    `--cel`, i nie przechodzi ono przez parser — więc to jedyna droga, którą tabulator dochodzi
    do komórki. Bez tego testu gałąź `\\t` z `FORMULA_PREFIXES` nie miałaby ani jednego
    wywołującego w eksporcie i byłaby drugą stałą deklarującą pokrycie, którego nie ma.
    """
    metadane: tuple[tuple[str, object], ...] = (("cel_pobrania", f"{prefiks}1+1"),)
    (sciezka,) = eksportuj(tmp_path / "m", zrodlo(wpis()), formaty=("xlsx",), metadane=metadane)

    assert ("cel_pobrania", f"'{prefiks}1+1") in arkusz(sciezka, "Metadane")


def test_bialy_prefiks_formuly_nie_dochodzi_do_kolumny_bo_parser_obcina_pole() -> None:
    """Antypustka podziału wyżej: parametryzacja z filtrem jest twierdzeniem o parserze.

    Gdyby `_napis` przestał obcinać białe znaki, tabulator wszedłby do kolumny metadanej drogą,
    której `PREFIKSY_WIDOCZNE` nie obejmuje — a filtr wyglądałby dalej tak samo. Druga asercja
    mówi przy okazji, że obcięcie nie odsłania formuły: `\\t=1+1` traci tabulator i zostaje
    zneutralizowane jako `=`.
    """
    assert set(FORMULA_PREFIXES) - set(PREFIKSY_WIDOCZNE) == {"\t"}
    assert wiersz(wpis(appellant_raw="\t=1+1"))["odwolujacy"] == "=1+1"
    assert sanitize_text(str(wiersz(wpis(appellant_raw="\t=1+1"))["odwolujacy"])) == "'=1+1"


def test_zapis_jest_atomowy_bez_pliku_tymczasowego_po_sukcesie(eksport: dict[str, Path]) -> None:
    for sciezka in eksport.values():
        assert not sciezka.with_name(f".{sciezka.name}.tmp").exists()
    assert not list(eksport["md"].glob(".*.tmp"))


def test_nieznany_format_jest_bledem_eksportu(tmp_path: Path) -> None:
    with pytest.raises(ExportError, match="zip"):
        eksportuj(tmp_path / "x", zrodlo(wpis()), formaty=("zip",), metadane=())


def test_sufiks_doklejany_do_nazwy_z_kropkami_w_rdzeniu(tmp_path: Path) -> None:
    """`with_suffix` uznałby `.2024-01-31` za rozszerzenie i obciął rdzeń."""
    rdzen = tmp_path / "kio_2024-01-01..2024-01-31"

    sciezki = exporter.sciezki_wyjsciowe(rdzen, FORMATY)

    assert sciezki["xlsx"].name == "kio_2024-01-01..2024-01-31.xlsx"
    assert sciezki["md"].name == "kio_2024-01-01..2024-01-31_md"


def test_nazwa_pliku_md_pochodzi_z_tozsamosci_nie_z_sygnatury() -> None:
    assert exporter.nazwa_pliku_md(wpis()) == "atlas_kio-1205-20.md"
